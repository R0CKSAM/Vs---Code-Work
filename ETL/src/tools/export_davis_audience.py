"""Export exact IP-level Davis Cup audience metrics to an Excel workbook."""
import argparse
import datetime as dt
import json
from pathlib import Path

import duckdb
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from build_concurrency import (
    DEFAULT_LAKE_FOLDER, PATH_MAP, channel_candidate_sql,
    configure_lake_selection, minute_ist_sql, parquet_source_sql,
)


def export(args):
    args.source, args.archive_lake = 'stream', []
    configure_lake_selection(args)
    start, end = dt.date.fromisoformat(args.start), dt.date.fromisoformat(args.end)
    dates = [(start + dt.timedelta(days=n)).isoformat() for n in range((end-start).days+1)]
    assert set(dates) == {p.date_text for p in args.selected_partitions}, 'Missing daily partitions'
    args.out.mkdir(parents=True, exist_ok=True)
    feeds = sorted((k, v) for k, v in PATH_MAP.items() if v.startswith('Davis Cup ('))
    assert len(feeds) == 6
    con = duckdb.connect()
    con.execute("SET threads=2")
    con.execute("SET memory_limit='3GB'")
    con.execute('SET temp_directory=?', [str(args.out/'scratch')])
    con.execute('CREATE TABLE feeds(asset VARCHAR, label VARCHAR)')
    con.executemany('INSERT INTO feeds VALUES (?,?)', feeds)
    con.execute('CREATE TABLE segments(minute TIMESTAMP, asset VARCHAR, ip VARCHAR, requests BIGINT)')
    coverage = []
    for part in args.selected_partitions:
        files = parquet_source_sql(list(part.files))
        # Retain non-Davis segments to establish exclusivity across the entire range.
        con.execute(f'''INSERT INTO segments
            SELECT minute, asset, ip, count(*) FROM (
                SELECT {minute_ist_sql()} AS minute,
                    {channel_candidate_sql('reqPath')} AS asset,
                    NULLIF(cliIP,'') AS ip
                FROM read_parquet({files}, hive_partitioning=true, union_by_name=true)
                WHERE lower(coalesce(reqPath,'')) LIKE '%.ts'
                  AND try_cast(reqTimeSec AS DOUBLE) IS NOT NULL
            ) WHERE CAST(minute AS DATE)=CAST(? AS DATE) GROUP BY 1,2,3''', [part.date_text])
        count = con.execute('SELECT sum(requests) FROM segments WHERE CAST(minute AS DATE)=CAST(? AS DATE)', [part.date_text]).fetchone()[0]
        coverage.append({'date':part.date_text,'ts_requests':count,'files':[str(p) for p in part.files]})
        print('Scanned',part.date_text,count,flush=True)
    con.execute('''CREATE TABLE exclusive_ips AS SELECT ip FROM segments s
        LEFT JOIN feeds f ON s.asset=f.asset WHERE ip IS NOT NULL
        GROUP BY ip HAVING count(*) FILTER(WHERE f.asset IS NOT NULL)>0
          AND count(*) FILTER(WHERE f.asset IS NULL)=0''')
    con.execute('''CREATE VIEW davis AS SELECT s.*,f.label,e.ip IS NOT NULL AS exclusive
        FROM segments s JOIN feeds f USING(asset) LEFT JOIN exclusive_ips e USING(ip)''')
    assert con.execute('''SELECT count(*) FROM segments s JOIN exclusive_ips e USING(ip)
        LEFT JOIN feeds f USING(asset) WHERE f.asset IS NULL''').fetchone()[0] == 0
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes='A2'
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='173D58')
        for i, header in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(i)].width=min(54,max(21,len(header)+2))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,float):
                    cell.number_format='0.000000'
        return ws

    notes=[
        ('Date range (IST)',f'{args.start} 00:00:00 to {args.end} 23:59:59'),
        ('Scope','STREAM only; six Davis Cup full asset IDs'),
        ('Watch hours','Estimated raw .ts requests * 6 / 3600, all HTTP statuses; retries included. Not measured player watch time.'),
        ('Concurrency','Exact distinct non-empty cliIP per IST minute with a .ts request. IP is not necessarily a person.'),
        ('Combined concurrency','Union of cliIP across all six feeds per minute, not the sum of individual counts.'),
        ('Exclusive cohort','IP with Davis Cup .ts requests and zero .ts requests on any other STREAM asset over the entire selected range. Unmapped assets count as other.'),
        ('Exclusive daily figures','Same whole-range exclusive cohort, measured on each day; daily distinct IPs must not be summed for the range total.'),
        ('Zero minutes','All 1440 minutes per day included. Zero means no qualifying requests in the available lake, not proof of collection completeness.'),
        ('Missing IP','Included in watch hours, excluded from distinct IPs and exclusive audience.'),
    ]
    ws=sheet('Read Me',['Item','Definition'],notes)
    ws.column_dimensions['B'].width=125
    measures='coalesce(sum(requests),0)/600.0, count(DISTINCT ip)'
    rows=[]
    for day in ['TOTAL']+dates:
        for asset,label in feeds+[(None,'ALL SIX COMBINED')]:
            where,params=['true'],[]
            if day!='TOTAL':
                where.append('CAST(minute AS DATE)=CAST(? AS DATE)');params.append(day)
            if asset:
                where.append('asset=?');params.append(asset)
            result=con.execute(f"SELECT {measures} FROM davis WHERE {' AND '.join(where)}",params).fetchone()
            rows.append((day,label,*result))
    sheet('Watch Hours',['Date IST','Feed','Raw watch hours estimated','Distinct cliIP'],rows)
    exclusive_rows=[]
    for day in ['TOTAL']+dates:
        where='exclusive'+(' AND CAST(minute AS DATE)=CAST(? AS DATE)' if day!='TOTAL' else '')
        result=con.execute(f'SELECT {measures} FROM davis WHERE {where}',[day] if day!='TOTAL' else []).fetchone()
        exclusive_rows.append((day,*result))
    sheet('Exclusive Audience',['Date IST','Raw watch hours estimated','Distinct exclusive cliIP'],exclusive_rows)
    minute_values={r[0]:r[1:] for r in con.execute('SELECT minute,count(DISTINCT ip),count(DISTINCT ip) FILTER(WHERE exclusive) FROM davis GROUP BY minute').fetchall()}
    per_feed={(r[0],r[1]):r[2] for r in con.execute('SELECT minute,asset,count(DISTINCT ip) FROM davis GROUP BY 1,2').fetchall()}
    minute_rows=[]
    for n in range(len(dates)*1440):
        minute=dt.datetime.combine(start,dt.time())+dt.timedelta(minutes=n)
        counts=[per_feed.get((minute,asset),0) for asset,_ in feeds]
        combined,exclusive=minute_values.get(minute,(0,0))
        assert max(counts)<=combined<=sum(counts)
        assert 0<=exclusive<=combined
        minute_rows.append((minute.strftime('%Y-%m-%d %H:%M:%S'),*counts,combined,exclusive))
    sheet('Minute Concurrency',['Minute IST']+[v for _,v in feeds]+['Combined distinct cliIP','Exclusive distinct cliIP'],minute_rows)
    sheet('Feed Links',['Feed','URL'],[(v,f'https://daviscup-veto.akamaized.net/{k}/playlist.m3u8') for k,v in feeds])
    sheet('Source Coverage',['Date IST','All STREAM TS requests','Source parquet files'],[(r['date'],r['ts_requests'],'; '.join(r['files'])) for r in coverage])
    total=sum(r[2] for r in rows[:6])
    assert abs(total-rows[6][2])<0.000001
    target=args.out/f'davis_cup_audience_{args.start}_to_{args.end}.xlsx'
    temp=target.with_suffix('.tmp.xlsx')
    wb.save(temp)
    check=load_workbook(temp,read_only=True,data_only=True)
    assert check['Minute Concurrency'].max_row==len(dates)*1440+1
    assert check['Watch Hours'].max_row==(len(dates)+1)*7+1
    check.close()
    temp.replace(target)
    audit={'dates':dates,'coverage':coverage,'feeds':feeds,'exclusive':exclusive_rows,'watch_rows':rows,'minute_rows':len(minute_rows),'checks_passed':True}
    target.with_suffix('.audit.json').write_text(json.dumps(audit,indent=2,default=str),encoding='utf-8')
    print('EXCLUSIVE',exclusive_rows,flush=True)
    print('WORKBOOK',target.resolve(),flush=True)
    con.close()


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--start',required=True)
    parser.add_argument('--end',required=True)
    parser.add_argument('--lake',type=Path,default=DEFAULT_LAKE_FOLDER)
    parser.add_argument('--out',type=Path,required=True)
    export(parser.parse_args())
