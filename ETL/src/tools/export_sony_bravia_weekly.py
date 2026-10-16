"""Read-only Sony BRAVIA lake comparison; exports aggregate counts, never identifiers."""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
import time
from pathlib import Path

import duckdb
import pandas as pd
import pyarrow as pa
import pyarrow.compute as pc
import pyarrow.parquet as pq

from build_concurrency import DEFAULT_LAKE_FOLDER, q
from build_identity_minute import normalized_identity_sql
from decode_distinct_ua_lookup import normalize_ua

ETL_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ETL_ROOT))
from src.common.lake_partitions import discover_partitions, resolve_lake_roots

IST = dt.timezone(dt.timedelta(hours=5, minutes=30))
PLACEHOLDERS = ('', 'null', 'none', 'undefined', 'unknown', 'nan', 'n/a', 'na', '0',
                '00000000-0000-0000-0000-000000000000')
MEASURES = """
    coalesce(sum(requests), 0)::BIGINT AS segment_requests,
    coalesce(sum(requests), 0) / 600.0 AS estimated_raw_watch_hours,
    coalesce(sum(ok), 0) / 600.0 AS estimated_http200_watch_hours,
    count(DISTINCT ip)::BIGINT AS distinct_client_ips,
    count(DISTINCT device)::BIGINT AS distinct_reported_device_ids,
    coalesce(sum(requests) FILTER(WHERE device IS NOT NULL), 0)::BIGINT AS requests_with_device_id,
    coalesce(sum(requests) FILTER(WHERE ip IS NOT NULL), 0)::BIGINT AS requests_with_ip,
    100.0 * coalesce(sum(requests) FILTER(WHERE device IS NOT NULL), 0)
        / nullif(sum(requests), 0) AS device_id_request_coverage_pct
"""


def is_bravia(ua: str, model: str = '') -> bool:
    return 'bravia' in normalize_ua(ua).casefold() or 'bravia' in str(model or '').casefold()


def pct_change(current: float, previous: float) -> float | None:
    return 100.0 * (current - previous) / previous if previous else None


def date_list(start: dt.date, end: dt.date) -> list[str]:
    return [(start + dt.timedelta(days=i)).isoformat() for i in range((end - start).days + 1)]


def verify_partitions(parts, sources, start, end):
    expected = {(source, day) for source in sources for day in date_list(start, end)}
    actual = {(part.source, part.date_text) for part in parts}
    if missing := expected - actual:
        raise ValueError(f'Missing lake partitions; refusing to interpret missing data as zero: {sorted(missing)}')


def build_aliases(con, lookup: Path, daily_ua: Path, start: str, end: str) -> int:
    decoded = pd.read_parquet(lookup, columns=['ua_norm', 'model']).fillna('')
    decoded = decoded.drop_duplicates('ua_norm', keep='last')
    raw = con.execute('SELECT DISTINCT userAgent FROM read_parquet(?) WHERE log_date BETWEEN ? AND ?',
                      [str(daily_ua), start, end]).fetchdf()
    raw['ua_norm'] = raw.userAgent.map(normalize_ua)
    joined = raw.merge(decoded, on='ua_norm', how='left', validate='many_to_one').fillna('')
    selected = joined[joined.apply(lambda row: is_bravia(row.ua_norm, row.model), axis=1)].copy()
    selected['model'] = selected.model.map(
        lambda model: ' '.join(model.upper().split()) if 'bravia' in model.casefold()
        else 'BRAVIA (MODEL UNSPECIFIED)')
    con.register('alias_frame', selected[['userAgent', 'model']].rename(columns={'userAgent': 'ua_raw'}))
    con.execute('CREATE TEMP TABLE bravia_alias AS SELECT * FROM alias_frame')
    con.unregister('alias_frame')
    return len(selected)


def initialize_matched(con):
    con.execute('''CREATE TEMP TABLE matched (
        source VARCHAR, day VARCHAR, model VARCHAR, ip VARCHAR, device VARCHAR,
        requests BIGINT, ok BIGINT, first_epoch DOUBLE, last_epoch DOUBLE)''')


def collect_partition(con, source: str, day: str, files: list[Path]):
    start = dt.datetime.combine(dt.date.fromisoformat(day), dt.time(), IST).timestamp()
    end = start + 86400
    device = normalized_identity_sql('device_id')
    invalid = ','.join("'" + value + "'" for value in PLACEHOLDERS)
    insert_sql = f'''
        INSERT INTO matched
        WITH candidates AS (
            SELECT UA, cliIP, queryStr, statusCode, TRY_CAST(reqTimeSec AS DOUBLE) AS epoch
            FROM bravia_batch
            WHERE lower(coalesce(reqPath, '')) LIKE '%.ts'
              AND TRY_CAST(reqTimeSec AS DOUBLE) >= ?
              AND TRY_CAST(reqTimeSec AS DOUBLE) < ?
        ), identities AS (
            SELECT *, {device} AS parsed_device FROM candidates
        )
        SELECT ?, ?, coalesce(a.model, 'BRAVIA (MODEL UNSPECIFIED)'),
               CASE WHEN lower(trim(coalesce(cliIP, ''))) IN ({invalid}) THEN NULL ELSE trim(cliIP) END,
               CASE WHEN lower(coalesce(parsed_device, '')) IN ({invalid}) THEN NULL ELSE parsed_device END,
               count(*), count(*) FILTER(WHERE TRY_CAST(statusCode AS INTEGER)=200),
               min(epoch), max(epoch)
        FROM identities r LEFT JOIN bravia_alias a ON r.UA=a.ua_raw
        GROUP BY 1,2,3,4,5
    '''
    aliases = pa.array([row[0] for row in con.execute('SELECT ua_raw FROM bravia_alias').fetchall()], type=pa.string())
    columns = ['cliIP', 'queryStr', 'statusCode', 'reqPath', 'reqTimeSec']
    scanned_groups, matched_groups = 0, 0
    # FAST query strings dominate compressed size. Read UA first and fetch the
    # other columns only for row groups with BRAVIA evidence, keeping RAM bounded.
    for file in files:
        reader = pq.ParquetFile(file)
        for group in range(reader.num_row_groups):
            ua = reader.read_row_group(group, columns=['UA'], use_threads=False).column('UA')
            mask = pc.or_kleene(pc.match_substring(ua, 'bravia', ignore_case=True),
                                pc.is_in(ua, value_set=aliases))
            selected = pc.indices_nonzero(pc.fill_null(mask, False))
            scanned_groups += 1
            if not len(selected):
                continue
            matched_groups += 1
            batch = reader.read_row_group(group, columns=columns, use_threads=True).take(selected)
            batch = batch.append_column('UA', ua.take(selected))
            con.register('bravia_batch', batch)
            try:
                con.execute(insert_sql, [start, end, source, day])
            finally:
                con.unregister('bravia_batch')
    return {'scanned_row_groups': scanned_groups, 'matched_row_groups': matched_groups}


def partition_metadata(part):
    bounds, rows = [], 0
    for path in part.files:
        meta = pq.ParquetFile(path).metadata
        rows += meta.num_rows
        for r in range(meta.num_row_groups):
            group = meta.row_group(r)
            for c in range(group.num_columns):
                col = group.column(c)
                if col.path_in_schema == 'reqTimeSec' and col.statistics and col.statistics.has_min_max:
                    bounds.extend([float(col.statistics.min), float(col.statistics.max)])
    first = dt.datetime.fromtimestamp(min(bounds), IST).isoformat() if bounds else None
    last = dt.datetime.fromtimestamp(max(bounds), IST).isoformat() if bounds else None
    return {'source': part.source, 'date_ist': part.date_text, 'lake_rows': rows,
            'files': [str(file) for file in part.files], 'first_request_ist': first,
            'last_request_ist': last}


def summaries(con, sources, previous_start, current_start, end):
    grid = pd.DataFrame([{'source': source, 'day': day} for source in sources
                         for day in date_list(previous_start, end)])
    con.register('date_grid', grid)
    daily = con.execute(f'''SELECT g.source, g.day AS date_ist, {MEASURES}
        FROM date_grid g LEFT JOIN matched m ON g.source=m.source AND g.day=m.day
        GROUP BY 1,2 ORDER BY 1,2''').fetchdf()
    con.execute('CREATE TEMP VIEW labeled AS SELECT *, CASE WHEN day >= '
                f"'{current_start.isoformat()}' THEN 'current' ELSE 'previous' END AS period FROM matched")
    periods = pd.DataFrame([{'source': s, 'period': p} for s in sources for p in ('previous', 'current')])
    con.register('period_grid', periods)
    weekly = con.execute(f'''SELECT g.source, g.period, {MEASURES}
        FROM period_grid g LEFT JOIN labeled m ON g.source=m.source AND g.period=m.period
        GROUP BY 1,2 ORDER BY 1,2''').fetchdf()
    models = con.execute(f'SELECT source, period, model, {MEASURES} FROM labeled GROUP BY 1,2,3 ORDER BY 1,2,4 DESC').fetchdf()
    changes = []
    for source in sources:
        part = weekly[weekly.source == source].set_index('period')
        for metric in ('estimated_raw_watch_hours', 'estimated_http200_watch_hours',
                       'distinct_client_ips', 'distinct_reported_device_ids'):
            prior, current = float(part.loc['previous', metric]), float(part.loc['current', metric])
            changes.append({'source': source, 'metric': metric, 'previous': prior, 'current': current,
                            'absolute_change': current - prior, 'change_pct': pct_change(current, prior)})
    daily_changes = []
    for source in sources:
        indexed = daily[daily.source == source].set_index('date_ist')
        for day in date_list(current_start, end):
            earlier = (dt.date.fromisoformat(day) - dt.timedelta(days=7)).isoformat()
            row = {'source': source, 'date_ist': day, 'previous_week_date': earlier}
            for metric in ('estimated_raw_watch_hours', 'distinct_client_ips', 'distinct_reported_device_ids'):
                prior, current = float(indexed.loc[earlier, metric]), float(indexed.loc[day, metric])
                row.update({metric: current, 'previous_' + metric: prior,
                            metric + '_change_pct': pct_change(current, prior)})
            daily_changes.append(row)
    return {'daily': daily, 'weekly': weekly, 'weekly_changes': pd.DataFrame(changes),
            'daily_vs_previous_week': pd.DataFrame(daily_changes), 'models': models}


def write_workbook(path, tables, notes):
    from openpyxl.styles import Font, PatternFill, Alignment
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name[:31], index=False)
        pd.DataFrame({'methodology': notes}).to_excel(writer, sheet_name='Read Me', index=False)
        for sheet in writer.book:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='155C63')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.row_dimensions[1].height = 42
            for column in sheet.columns:
                values = [len(str(cell.value or '')) for cell in column]
                sheet.column_dimensions[column[0].column_letter].width = min(44, max(18, max(values) + 2))
                for cell in column[1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = '#,##0.00'
                    if sheet.title == 'Read Me':
                        cell.alignment = Alignment(wrap_text=True, vertical='top')
            if sheet.title == 'Read Me':
                sheet.column_dimensions['A'].width = 110
                for row in range(2, sheet.max_row + 1):
                    sheet.row_dimensions[row].height = 42


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--end', required=True, type=dt.date.fromisoformat)
    parser.add_argument('--sources', nargs='+', default=['stream', 'fast'], choices=['stream', 'fast'])
    parser.add_argument('--lake', type=Path, default=DEFAULT_LAKE_FOLDER)
    parser.add_argument('--archive-lake', type=Path, action='append')
    parser.add_argument('--out-dir', type=Path, required=True)
    parser.add_argument('--threads', type=int, default=4)
    args = parser.parse_args()
    if args.end >= dt.datetime.now(IST).date():
        raise SystemExit('Choose a completed past day, not today or a future date.')
    current_start = args.end - dt.timedelta(days=6)
    previous_start = current_start - dt.timedelta(days=7)
    roots = resolve_lake_roots(args.lake, args.archive_lake)
    parts = discover_partitions(roots, sources=args.sources, start=previous_start, end=args.end)
    verify_partitions(parts, args.sources, previous_start, args.end)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    pa.set_cpu_count(max(1, args.threads))
    con = duckdb.connect()
    con.execute(f'SET threads={max(1, args.threads)}')
    con.execute("SET memory_limit='2GB'")
    con.execute('SET preserve_insertion_order=false')
    con.execute(f"SET temp_directory='{q(args.out_dir / 'scratch')}'")
    lookup = ETL_ROOT / 'output/device_decode/ua_decode_lookup_both_all.parquet'
    lookup_modified = lookup.stat().st_mtime
    daily_ua = ETL_ROOT / 'output/watch_hours/daily_tables/user_agents_daily.parquet'
    aliases = build_aliases(con, lookup, daily_ua, previous_start.isoformat(), args.end.isoformat())
    initialize_matched(con)
    coverage = []
    print(f'Comparing {current_start} to {args.end} vs {previous_start} to {current_start-dt.timedelta(days=1)}; aliases={aliases}', flush=True)
    started = time.monotonic()
    for index, part in enumerate(parts, 1):
        step = time.monotonic()
        meta = partition_metadata(part)
        meta.update(collect_partition(con, part.source, part.date_text, list(part.files)))
        result = con.execute(f'SELECT {MEASURES} FROM matched WHERE source=? AND day=?',
                             [part.source, part.date_text]).fetchdf().iloc[0].to_dict()
        coverage.append(meta | result)
        progress = {'complete_partitions': index, 'total_partitions': len(parts), 'latest': meta | result,
                    'elapsed_seconds': round(time.monotonic()-started, 1)}
        (args.out_dir / 'progress.json').write_text(json.dumps(progress, indent=2), encoding='utf-8')
        print(f"{index}/{len(parts)} {part.key}: {result['segment_requests']:,.0f} segments, "
              f"{result['distinct_client_ips']:,.0f} IPs, {result['distinct_reported_device_ids']:,.0f} IDs "
              f"({time.monotonic()-step:.1f}s)", flush=True)
    tables = summaries(con, args.sources, previous_start, current_start, args.end)
    notes = [
        f'Current week: {current_start} to {args.end}; previous week: {previous_start} to {current_start-dt.timedelta(days=1)}. All dates IST.',
        'STREAM and FAST are reported separately, across all channels. Selected complete date partitions are read from the lake, not the UA preview HTML.',
        'Cohort: explicit BRAVIA in normalized UA or decoded model. Sony smartphones, PlayStation and generic Sony without BRAVIA evidence are excluded.',
        'Estimated raw watch hours = all matching .ts requests * 6 seconds / 3600. This includes retries/errors and is not measured player playback time.',
        'HTTP-200 watch hours use the same 6-second assumption for status 200 requests only. Matching timestamps are constrained to the IST reporting day.',
        'Distinct client IPs are exact within each reporting group, not people or physical TVs. Shared IPs and IP changes affect this proxy.',
        'Distinct reported device IDs exclude common blank/placeholder values. Device-ID coverage is shown; zero means no observed usable IDs, not zero televisions.',
        'Weekly counts are recomputed distinct unions, never sums of daily counts. Model counts can overlap; do not sum them into a total device count.',
        'Week-on-week percentages with a zero previous baseline are blank, not a fabricated percent. Model unspecified means BRAVIA evidence exists but no reliable model label.',
        'No raw IP addresses or device IDs are included in these exports. UA identity is inferred from request metadata; no unique physical-TV count is claimed.',
    ]
    for name, table in tables.items():
        table.to_csv(args.out_dir / f'sony_bravia_{name}.csv', index=False, encoding='utf-8-sig')
    write_workbook(args.out_dir / 'sony_bravia_weekly_comparison.xlsx', tables, notes)
    manifest = {'generated_at_ist': dt.datetime.now(IST).isoformat(), 'current_start': str(current_start),
                'end': str(args.end), 'previous_start': str(previous_start), 'sources': args.sources,
                'lookup': str(lookup), 'lookup_modified_epoch': lookup_modified,
                'aliases': aliases, 'notes': notes, 'coverage': coverage,
                'elapsed_seconds': round(time.monotonic()-started, 1)}
    (args.out_dir / 'manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
    (args.out_dir / 'README.txt').write_text('\n\n'.join(notes)+'\n', encoding='utf-8')
    print(tables['weekly_changes'].to_string(index=False), flush=True)
    con.close()


if __name__ == '__main__':
    main()
