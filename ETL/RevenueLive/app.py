"""RevenueLive: channel-scoped revenue reporting and audited spreadsheet imports."""
import argparse
import csv
import datetime as dt
import functools
import hashlib
import io
import json
import os
from pathlib import Path
import secrets
import sqlite3
import threading
import time
from contextlib import closing
from decimal import Decimal, InvalidOperation

from flask import Flask, g, jsonify, request, send_from_directory
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

ROOT = Path(__file__).resolve().parent
HEADERS = ['Date', 'Channel Name', 'Views', 'Ad Impressions', 'Ad Revenue', 'Sponsorship/Others', 'Total Revenue']


class InvalidData(ValueError):
    pass


def parse_upload(content, suffix):
    if suffix == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        rows = []
        for i in range(sheet.nrows):
            row = sheet.row_values(i)
            if i and sheet.cell_type(i, 0) == xlrd.XL_CELL_DATE:
                row[0] = xlrd.xldate_as_datetime(row[0], book.datemode)
            rows.append(row)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        import zipfile
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100_000_000:
                raise InvalidData('Workbook expands beyond the 100 MB limit.')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            rows = list(book.worksheets[0].iter_rows(values_only=True))
        finally:
            book.close()
    elif suffix == '.csv':
        rows = list(csv.reader(io.StringIO(content.decode('utf-8-sig'))))
    else:
        raise InvalidData('Upload an XLS, XLSX or UTF-8 CSV file.')
    if not rows or [str(x or '').strip() for x in rows[0]] != HEADERS:
        raise InvalidData('Columns must match: ' + ', '.join(HEADERS))
    if len(rows) > 20001:
        raise InvalidData('Maximum 20,000 data rows per upload.')
    result, seen = [], set()
    for number, row in enumerate(rows[1:], 2):
        if all(x is None or str(x).strip() == '' for x in row):
            continue
        if len(row) != 7:
            raise InvalidData(f'Row {number}: expected seven columns.')
        date = row[0]
        if isinstance(date, dt.datetime):
            date = date.date()
        elif not isinstance(date, dt.date):
            try:
                date = dt.date.fromisoformat(str(date).strip())
            except ValueError:
                raise InvalidData(f'Row {number}: use an Excel date or YYYY-MM-DD.')
        channel = str(row[1] or '').strip()
        if not channel or len(channel) > 120:
            raise InvalidData(f'Row {number}: invalid channel name.')
        key = (date.isoformat(), channel.casefold())
        if key in seen:
            raise InvalidData(f'Row {number}: duplicate date/channel inside the file.')
        seen.add(key)
        values = []
        for index, value in enumerate(row[2:], 2):
            try:
                numeric = Decimal(str(value).strip())
                if not numeric.is_finite() or numeric < 0 or numeric > Decimal('1000000000000'):
                    raise ValueError()
                scaled = numeric * (100 if index >= 4 else 1)
                if scaled != scaled.to_integral_value():
                    raise ValueError()
                values.append(int(scaled))
            except (InvalidOperation, ValueError):
                raise InvalidData(f'Row {number}: {HEADERS[index]} must be non-negative; counts are integers and INR allows two decimals.')
        if values[2] + values[3] != values[4]:
            raise InvalidData(f'Row {number}: total revenue must equal ad revenue plus sponsorship/others.')
        result.append(dict(day=date.isoformat(), channel=channel, views=values[0], impressions=values[1], ad=values[2], other=values[3], total=values[4]))
    if not result:
        raise InvalidData('No data rows found.')
    return result


def create_app(data_dir=None):
    app = Flask(__name__, static_folder='static')
    data = Path(data_dir or os.environ.get('REVENUE_DATA_DIR', ROOT / 'data'))
    data.mkdir(parents=True, exist_ok=True)
    (data / 'uploads').mkdir(exist_ok=True)
    app.config.update(DATA_DIR=data, MAX_CONTENT_LENGTH=10*1024*1024)
    cookie_name='revenue_'+hashlib.sha256(str(data.resolve()).encode()).hexdigest()[:12]

    def db():
        if 'db' not in g:
            g.db = sqlite3.connect(data / 'revenuelive.db', timeout=30)
            g.db.row_factory = sqlite3.Row
            g.db.execute('PRAGMA foreign_keys=ON')
        return g.db

    @app.teardown_appcontext
    def close(_error):
        connection = g.pop('db', None)
        if connection:
            connection.close()

    with app.app_context():
        db().executescript('''
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE COLLATE NOCASE, password TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('admin','uploader','viewer')), active INTEGER NOT NULL DEFAULT 1, must_change INTEGER NOT NULL DEFAULT 0);
            CREATE TABLE IF NOT EXISTS channels(id INTEGER PRIMARY KEY, name TEXT UNIQUE COLLATE NOCASE NOT NULL);
            CREATE TABLE IF NOT EXISTS assignments(user_id INTEGER REFERENCES users(id), channel_id INTEGER REFERENCES channels(id), PRIMARY KEY(user_id,channel_id));
            CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY, user_id INTEGER REFERENCES users(id), csrf TEXT NOT NULL, expires REAL NOT NULL);
            CREATE TABLE IF NOT EXISTS attempts(ip TEXT PRIMARY KEY, failures INTEGER, expires REAL);
            CREATE TABLE IF NOT EXISTS uploads(id TEXT PRIMARY KEY, user_id INTEGER REFERENCES users(id), filename TEXT, digest TEXT, created TEXT, state TEXT, rows_json TEXT, preview_json TEXT);
            CREATE TABLE IF NOT EXISTS records(day TEXT, channel_id INTEGER REFERENCES channels(id), views INTEGER, impressions INTEGER, ad INTEGER, other INTEGER, total INTEGER, upload_id TEXT, PRIMARY KEY(day,channel_id));
            CREATE TABLE IF NOT EXISTS revisions(upload_id TEXT, day TEXT, channel_id INTEGER, previous TEXT);
            CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, created TEXT, user_id INTEGER, action TEXT, detail TEXT);
            CREATE TABLE IF NOT EXISTS super_admin(singleton INTEGER PRIMARY KEY CHECK(singleton=1), user_id INTEGER UNIQUE NOT NULL REFERENCES users(id));
        ''')
        db().commit()

    def log(action, detail=''):
        db().execute('INSERT INTO audit(created,user_id,action,detail) VALUES (?,?,?,?)', (dt.datetime.now(dt.timezone.utc).isoformat(), g.user['id'], action, detail))

    def permitted():
        if g.user['role'] == 'admin':
            return db().execute('SELECT * FROM channels ORDER BY name COLLATE NOCASE').fetchall()
        return db().execute('SELECT c.* FROM channels c JOIN assignments a ON c.id=a.channel_id WHERE a.user_id=? ORDER BY c.name COLLATE NOCASE', (g.user['id'],)).fetchall()

    def require(*roles):
        def decorator(fn):
            @functools.wraps(fn)
            def wrapped(*args, **kwargs):
                if not getattr(g, 'user', None):
                    return jsonify(error='Please sign in.'), 401
                if g.user['must_change'] and request.path not in {'/api/me','/api/password','/api/logout'}:
                    return jsonify(error='Change your temporary password first.'), 403
                if roles and g.user['role'] not in roles:
                    return jsonify(error='You do not have permission for this action.'), 403
                return fn(*args, **kwargs)
            return wrapped
        return decorator

    @app.before_request
    def auth():
        if not request.path.startswith('/api/'):
            return
        g.user = None
        raw = request.cookies.get(cookie_name, '')
        digest = hashlib.sha256(raw.encode()).hexdigest()
        session = db().execute('SELECT s.*,u.username,u.role,u.active,u.must_change FROM sessions s JOIN users u ON u.id=s.user_id WHERE token=? AND expires>? AND active=1', (digest,time.time())).fetchone()
        if session:
            g.user = dict(id=session['user_id'], username=session['username'], role=session['role'], must_change=session['must_change'])
            g.user['super_admin'] = bool(db().execute('SELECT 1 FROM super_admin WHERE user_id=?',(session['user_id'],)).fetchone())
            g.session = session
        if request.method in {'POST','PUT','DELETE'}:
            # Reject cross-site writes even on login, where no session exists yet.
            origin = request.headers.get('Origin')
            if origin and origin != request.host_url.rstrip('/'):
                return jsonify(error='Cross-origin requests are not allowed.'), 403
            if request.path not in {'/api/login','/api/account/request','/api/account/complete'} and (not session or not secrets.compare_digest(request.headers.get('X-CSRF-Token',''),session['csrf'])):
                return jsonify(error='Session expired. Sign in again.'), 403

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['X-Frame-Options']='DENY'
        response.headers['Referrer-Policy']='same-origin'
        response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        if request.path.startswith('/api/'):
            response.headers['Cache-Control']='no-store'
        return response

    @app.errorhandler(413)
    def large(_):
        return jsonify(error='Maximum file size is 10 MB.'),413

    @app.errorhandler(InvalidData)
    def invalid(error):
        return jsonify(error=str(error)),400

    @app.get('/')
    def index():
        return send_from_directory(ROOT / 'static','index.html')

    @app.get('/health')
    def health():
        return jsonify(service='revenuelive', status='ok')

    @app.post('/api/login')
    def login():
        value = request.get_json(silent=True) or {}
        ip = request.remote_addr
        attempt = db().execute('SELECT * FROM attempts WHERE ip=?',(ip,)).fetchone()
        if attempt and attempt['expires']>time.time() and attempt['failures']>=10:
            return jsonify(error='Too many attempts. Try again in 15 minutes.'),429
        user = db().execute('SELECT * FROM users WHERE username=? AND active=1',(str(value.get('username','')).strip()[:100],)).fetchone()
        password = str(value.get('password',''))
        if not user or len(password)>256 or not check_password_hash(user['password'], password):
            failures = attempt['failures']+1 if attempt and attempt['expires']>time.time() else 1
            db().execute('INSERT OR REPLACE INTO attempts VALUES (?,?,?)',(ip,failures,time.time()+900))
            db().commit()
            return jsonify(error='Invalid username or password.'),401
        raw, csrf = secrets.token_urlsafe(32), secrets.token_urlsafe(32)
        db().execute('DELETE FROM attempts WHERE ip=?',(ip,))
        db().execute('DELETE FROM sessions WHERE expires<?',(time.time(),))
        db().execute('INSERT INTO sessions VALUES (?,?,?,?)',(hashlib.sha256(raw.encode()).hexdigest(),user['id'],csrf,time.time()+8*3600))
        db().commit()
        response=jsonify(ok=True)
        response.set_cookie(cookie_name,raw,httponly=True,samesite='Strict',secure=os.environ.get('REVENUE_HTTPS')=='1',max_age=8*3600)
        return response

    @app.get('/api/me')
    @require()
    def me():
        return jsonify(user=g.user,csrf=g.session['csrf'],channels=[dict(c) for c in permitted()],demo=(data/'DEMO_DATA.json').exists())

    @app.post('/api/logout')
    @require()
    def logout():
        db().execute('DELETE FROM sessions WHERE token=?',(g.session['token'],));db().commit()
        response=jsonify(ok=True);response.delete_cookie(cookie_name);return response

    @app.post('/api/password')
    @require()
    def password():
        body=request.get_json() or {}
        value=str(body.get('password',''))
        if not 12<=len(value)<=256:
            raise InvalidData('Password must be 12 to 256 characters.')
        current=db().execute('SELECT password FROM users WHERE id=?',(g.user['id'],)).fetchone()[0]
        if not check_password_hash(current,str(body.get('current',''))):
            raise InvalidData('Current password is incorrect.')
        db().execute('UPDATE users SET password=?,must_change=0 WHERE id=?',(generate_password_hash(value),g.user['id']))
        db().execute('DELETE FROM email_tokens WHERE user_id=?',(g.user['id'],))
        db().execute('DELETE FROM sessions WHERE user_id=? AND token<>?',(g.user['id'],g.session['token']))
        log('password_changed');db().commit();return jsonify(ok=True)

    def report_rows():
        ids={c['id'] for c in permitted()}
        selected=[v for v in request.args.getlist('channel') if v]
        if selected == ['none']:
            ids=set()
        elif selected:
            try:
                chosen={int(value) for value in selected}
            except ValueError:
                raise InvalidData('Invalid channel.')
            if not chosen <= ids:
                raise InvalidData('Channel is not assigned to your account.')
            ids=chosen
        for key in ('start','end'):
            if request.args.get(key):
                try:
                    dt.date.fromisoformat(request.args[key])
                except ValueError:
                    raise InvalidData('Use YYYY-MM-DD dates.')
        start,end=request.args.get('start',''),request.args.get('end','9999-12-31') or '9999-12-31'
        if start>end:
            raise InvalidData('Start date must not be after end date.')
        placeholders=','.join('?' for _ in ids) or 'NULL'
        return [dict(r) for r in db().execute(f'SELECT r.*,c.name AS channel FROM records r JOIN channels c ON c.id=r.channel_id WHERE r.channel_id IN ({placeholders}) AND day>=? AND day<=? ORDER BY day DESC,c.name COLLATE NOCASE',(*sorted(ids),start,end))]

    @app.get('/api/report')
    @require()
    def report():
        rows=report_rows()
        totals={key:sum(r[key] for r in rows) for key in ('views','impressions','ad','other','total')}
        return jsonify(rows=rows,totals=totals,currency='INR',money_unit='paise')

    @app.get('/api/export')
    @require()
    def download():
        buffer=io.StringIO();writer=csv.writer(buffer);writer.writerow(HEADERS)
        for r in report_rows():
            channel=r['channel']
            if channel.startswith(('=','+','-','@')):
                channel="'"+channel
            writer.writerow([r['day'],channel,r['views'],r['impressions'],*[f'{r[k]/100:.2f}' for k in ('ad','other','total')]])
        response=app.response_class('\ufeff'+buffer.getvalue(),mimetype='text/csv')
        filename='DEMO_revenue.csv' if (data/'DEMO_DATA.json').exists() else 'revenue.csv'
        response.headers['Content-Disposition']=f'attachment; filename={filename}'
        return response

    def resolve_rows(rows):
        allowed={c['id'] for c in permitted()}
        names={c['name'].casefold():c['id'] for c in db().execute('SELECT * FROM channels')}
        resolved=[]
        for row in rows:
            cid=names.get(row['channel'].casefold())
            if cid is None or cid not in allowed:
                raise InvalidData('File contains unknown or unassigned channels. Ask the admin to register/assign all channels before uploading.')
            resolved.append({**row,'channel_id':cid})
        return resolved

    @app.post('/api/uploads/preview')
    @require('admin','uploader')
    def preview():
        incoming=request.files.get('file')
        if not incoming or not incoming.filename:
            raise InvalidData('Choose a file.')
        content=incoming.read()
        try:
            rows=parse_upload(content,Path(incoming.filename).suffix.lower())
        except InvalidData:
            raise
        except Exception:
            raise InvalidData('Unable to read workbook. Check its format and date cells.')
        rows=resolve_rows(rows)
        digest=hashlib.sha256(content).hexdigest()
        if db().execute("SELECT 1 FROM uploads WHERE digest=? AND state='committed'",(digest,)).fetchone():
            raise InvalidData('This exact file has already been published.')
        originals=[]
        for row in rows:
            old=db().execute('SELECT * FROM records WHERE day=? AND channel_id=?',(row['day'],row['channel_id'])).fetchone()
            originals.append(dict(old) if old else None)
        uid=secrets.token_hex(16)
        filename=secure_filename(incoming.filename) or 'upload'
        (data/'uploads'/f'{uid}_{filename}').write_bytes(content)
        db().execute('INSERT INTO uploads VALUES (?,?,?,?,?,?,?,?)',(uid,g.user['id'],filename,digest,dt.datetime.now(dt.timezone.utc).isoformat(),'pending',json.dumps(rows),json.dumps(originals)))
        db().commit()
        return jsonify(id=uid,rows=rows,duplicates=sum(x is not None for x in originals))

    @app.post('/api/uploads/<uid>/commit')
    @require('admin','uploader')
    def commit(uid):
        db().execute('BEGIN IMMEDIATE')
        upload=db().execute('SELECT * FROM uploads WHERE id=? AND user_id=?',(uid,g.user['id'])).fetchone()
        if not upload or upload['state']!='pending':
            raise InvalidData('Upload is unavailable or already published.')
        rows=resolve_rows(json.loads(upload['rows_json']))
        original=json.loads(upload['preview_json'])
        if any(x is not None for x in original) and not (request.get_json(silent=True) or {}).get('replace') is True:
            raise InvalidData('Confirm replacement of existing date/channel records.')
        for row,old in zip(rows,original):
            current=db().execute('SELECT * FROM records WHERE day=? AND channel_id=?',(row['day'],row['channel_id'])).fetchone()
            if (dict(current) if current else None)!=old:
                raise InvalidData('Data changed after preview. Upload again to review the latest version.')
            db().execute('INSERT INTO revisions VALUES (?,?,?,?)',(uid,row['day'],row['channel_id'],json.dumps(old)))
            db().execute('INSERT OR REPLACE INTO records VALUES (?,?,?,?,?,?,?,?)',(row['day'],row['channel_id'],row['views'],row['impressions'],row['ad'],row['other'],row['total'],uid))
        db().execute("UPDATE uploads SET state='committed' WHERE id=?",(uid,))
        log('upload_published',uid);db().commit();return jsonify(ok=True)

    @app.get('/api/uploads')
    @require('admin','uploader')
    def uploads():
        where='' if g.user['role']=='admin' else ' WHERE u.user_id=?'
        args=() if not where else (g.user['id'],)
        rows=db().execute('SELECT u.id,u.filename,u.created,u.state,v.username FROM uploads u JOIN users v ON v.id=u.user_id'+where+' ORDER BY u.created DESC LIMIT 100',args).fetchall()
        return jsonify(rows=[dict(r) for r in rows])

    @app.post('/api/uploads/<uid>/restore')
    @require('admin')
    def restore(uid):
        db().execute('BEGIN IMMEDIATE')
        upload=db().execute("SELECT * FROM uploads WHERE id=? AND state='committed'",(uid,)).fetchone()
        if not upload:
            raise InvalidData('Only a published upload can be rolled back.')
        revisions=db().execute('SELECT * FROM revisions WHERE upload_id=?',(uid,)).fetchall()
        for rev in revisions:
            current=db().execute('SELECT upload_id FROM records WHERE day=? AND channel_id=?',(rev['day'],rev['channel_id'])).fetchone()
            if not current or current[0]!=uid:
                raise InvalidData('A newer upload changed these records. Roll back the newer upload first.')
        for rev in revisions:
            old=json.loads(rev['previous'])
            db().execute('DELETE FROM records WHERE day=? AND channel_id=?',(rev['day'],rev['channel_id']))
            if old:
                db().execute('INSERT INTO records VALUES (?,?,?,?,?,?,?,?)',tuple(old[k] for k in ('day','channel_id','views','impressions','ad','other','total','upload_id')))
        db().execute("UPDATE uploads SET state='restored' WHERE id=?",(uid,))
        log('upload_rolled_back',uid);db().commit();return jsonify(ok=True)

    @app.get('/api/admin/users')
    @require('admin')
    def users():
        rows=[]
        for user in db().execute('SELECT id,username,role,active,must_change FROM users ORDER BY username'):
            rows.append({**dict(user),'super_admin':bool(db().execute('SELECT 1 FROM super_admin WHERE user_id=?',(user['id'],)).fetchone()),'channels':[r[0] for r in db().execute('SELECT channel_id FROM assignments WHERE user_id=?',(user['id'],))]})
        return jsonify(users=rows,channels=[dict(c) for c in permitted()])

    @app.post('/api/admin/channels')
    @require('admin')
    def channels():
        name=str((request.get_json() or {}).get('name','')).strip()
        if not name or len(name)>120:
            raise InvalidData('Enter a channel name, up to 120 characters.')
        try:
            db().execute('INSERT INTO channels(name) VALUES (?)',(name,))
        except sqlite3.IntegrityError:
            raise InvalidData('Channel already exists.')
        log('channel_created',name);db().commit();return jsonify(ok=True)

    from account_email import install
    email_address, mail_settings, send_invitation = install(app, db, data, InvalidData)

    @app.post('/api/admin/users')
    @require('admin')
    def save_user():
        body=request.get_json() or {}
        username=str(body.get('username','')).strip()
        role=body.get('role')
        password=str(body.get('password',''))
        uid=body.get('id')
        invite = body.get('invite') is True
        if invite:
            if uid:
                raise InvalidData('Invitations are for new accounts. Existing email users can use Forgot password.')
            username = email_address(username)
            mail_settings()
            password = secrets.token_urlsafe(48)
        if uid is not None and (type(uid)!=int or uid<=0):
            raise InvalidData('Invalid user ID.')
        if role not in {'admin','uploader','viewer'} or not username or len(username)>80:
            raise InvalidData('Enter a username and valid role.')
        if (not uid or password) and not 12<=len(password)<=256:
            raise InvalidData('Temporary password must be 12 to 256 characters.')
        ids=body.get('channels',[])
        if not isinstance(ids,list) or any(type(x)!=int for x in ids):
            raise InvalidData('Invalid channel assignments.')
        all_ids={r[0] for r in db().execute('SELECT id FROM channels')}
        if not set(ids)<=all_ids:
            raise InvalidData('Unknown channel assignment.')
        active=int(body.get('active',True) is True)
        if uid==g.user['id'] and (role!='admin' or not active):
            raise InvalidData('You cannot remove your own admin access.')
        try:
            db().execute('BEGIN IMMEDIATE')
            owner=db().execute('SELECT user_id FROM super_admin WHERE singleton=1').fetchone()
            existing=db().execute('SELECT role FROM users WHERE id=?',(uid,)).fetchone() if uid else None
            if owner and uid==owner['user_id']:
                raise InvalidData('The Super Admin account is protected. Use Change password for your own password.')
            if not owner or owner['user_id']!=g.user['id']:
                if role=='admin' or (existing and existing['role']=='admin'):
                    return jsonify(error='Only the Super Admin can create or modify admin accounts.'),403
            if uid:
                if not db().execute('SELECT 1 FROM users WHERE id=?',(uid,)).fetchone():
                    raise InvalidData('User not found.')
                email_account=db().execute('SELECT email FROM email_accounts WHERE user_id=?',(uid,)).fetchone()
                if email_account and username.casefold()!=email_account['email'].casefold():
                    raise InvalidData('Email login cannot be renamed. Disable this account and invite the new email separately.')
                db().execute('UPDATE users SET username=?,role=?,active=? WHERE id=?',(username,role,active,uid))
                if password:
                    db().execute('UPDATE users SET password=?,must_change=1 WHERE id=?',(generate_password_hash(password),uid))
                if password or not active:
                    db().execute('DELETE FROM sessions WHERE user_id=? AND token<>?',(uid,g.session['token']))
                    db().execute('DELETE FROM email_tokens WHERE user_id=?',(uid,))
            else:
                uid=db().execute('INSERT INTO users(username,password,role,active,must_change) VALUES (?,?,?,?,1)',(username,generate_password_hash(password),role,active)).lastrowid
            db().execute('DELETE FROM assignments WHERE user_id=?',(uid,))
            db().executemany('INSERT INTO assignments VALUES (?,?)',[(uid,c) for c in set(ids)])
            if invite:
                db().execute('INSERT INTO email_accounts(user_id,email) VALUES (?,?)',(uid,username))
            log('user_saved',str(uid));db().commit()
        except sqlite3.IntegrityError:
            raise InvalidData('Username already exists.')
        if invite:
            send_invitation(uid)
        return jsonify(ok=True)

    return app


def bootstrap(data):
    app=create_app(data)
    with closing(sqlite3.connect(data/'revenuelive.db')) as con, con:
        if not con.execute('SELECT 1 FROM users LIMIT 1').fetchone():
            password=secrets.token_urlsafe(18)
            con.execute("INSERT INTO users(username,password,role,must_change) VALUES (?,?,'admin',1)",('admin',generate_password_hash(password)))
            (data/'initial_admin.txt').write_text('Username: admin\nTemporary password: '+password+'\nChange at first login. Keep this file private and delete after changing.\n',encoding='utf-8')
        sample=ROOT/'Upload File.xls'
        if sample.exists() and not con.execute('SELECT 1 FROM channels LIMIT 1').fetchone():
            for row in parse_upload(sample.read_bytes(),'.xls'):
                con.execute('INSERT OR IGNORE INTO channels(name) VALUES (?)',(row['channel'],))
    return app


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--host',default='127.0.0.1')
    parser.add_argument('--port',type=int,default=8820)
    parser.add_argument('--demo',action='store_true')
    args=parser.parse_args()
    data=ROOT/'demo_data' if args.demo else Path(os.environ.get('REVENUE_DATA_DIR',ROOT/'data'))
    app=bootstrap(data)
    if args.demo:
        from demo_data import seed_demo
        seed_demo(data)
    from waitress import create_server
    server=create_server(app,host=args.host,port=args.port,threads=4)
    stop=data/f'stop_{args.port}'
    stop.unlink(missing_ok=True)
    def monitor():
        while not stop.exists():
            time.sleep(1)
        server.close()
        os._exit(0)
    threading.Thread(target=monitor,daemon=True).start()
    print(f'RevenueLive: http://{args.host}:{args.port}',flush=True)
    server.run()
