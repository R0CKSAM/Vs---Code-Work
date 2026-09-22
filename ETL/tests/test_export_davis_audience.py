import argparse
import datetime as dt
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import duckdb
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src/tools'))
import export_davis_audience as exporter


def test_exact_union_and_cross_day_exclusivity(tmp_path, monkeypatch):
    feeds=sorted(k for k,v in exporter.PATH_MAP.items() if v.startswith('Davis Cup ('))
    start=dt.datetime(2026,9,17,tzinfo=dt.timezone(dt.timedelta(hours=5,minutes=30)))
    partitions=[]
    with duckdb.connect() as con:
        con.execute('CREATE TABLE raw(reqTimeSec DOUBLE,reqPath VARCHAR,cliIP VARCHAR)')
        for n in range(3):
            day=start+dt.timedelta(days=n)
            rows=[]
            if n==0:
                rows=[(day.timestamp(),f'/{feeds[0]}/s.ts','shared'),
                      (day.timestamp(),f'/{feeds[1]}/s.ts','shared'),
                      (day.timestamp(),f'/{feeds[0]}/s.ts','exclusive'),
                      (day.timestamp(),f'/{feeds[0]}/s.ts',None)]
            if n==1:
                rows=[(day.timestamp(),'/other/s.ts','shared')]
            con.execute('DELETE FROM raw')
            if rows:
                con.executemany('INSERT INTO raw VALUES (?,?,?)',rows)
            file=tmp_path/f'{n}.parquet'
            con.execute('COPY raw TO ? (FORMAT PARQUET)',[str(file)])
            partitions.append(SimpleNamespace(date_text=day.date().isoformat(),files=[file]))
    def configure(args):
        args.selected_partitions=partitions
    monkeypatch.setattr(exporter,'configure_lake_selection',configure)
    args=argparse.Namespace(start='2026-09-17',end='2026-09-19',out=tmp_path/'out',lake=tmp_path)
    exporter.export(args)
    book=load_workbook(next(args.out.glob('*.xlsx')),read_only=True,data_only=True)
    minute=list(book['Minute Concurrency'].values)[1]
    assert minute[1:3]==(2,1)
    assert minute[-2:]==(2,1)
    exclusive=list(book['Exclusive Audience'].values)[1]
    assert exclusive[2]==1
    assert abs(exclusive[1]-1/600)<1e-12
    assert book['Minute Concurrency'].max_row==4321
    book.close()
    audit=json.loads(next(args.out.glob('*.audit.json')).read_text())
    assert audit['checks_passed']
