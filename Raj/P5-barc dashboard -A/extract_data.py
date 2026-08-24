import os
import sys
import subprocess
import json

def install_and_import(package, import_name=None):
    if import_name is None:
        import_name = package
    try:
        __import__(import_name)
    except ImportError:
        print(f"Installing {package}...")
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", package])
        except Exception as e:
            print(f"Failed to install {package} via pip: {e}")
            print("Please run: pip install python-pptx pandas openpyxl")
            sys.exit(1)

# Ensure required libraries are installed
install_and_import("python-pptx", "pptx")
install_and_import("pandas")
install_and_import("openpyxl")

import pandas as pd
from pptx import Presentation

def extract_pptx(pptx_path, out_txt_path):
    print(f"Extracting PPTX: {pptx_path}...")
    try:
        prs = Presentation(pptx_path)
        content = []
        for i, slide in enumerate(prs.slides):
            slide_text = [f"=== Slide {i+1} ==="]
            
            # Extract slide title if available
            if slide.shapes.title:
                slide_text.append(f"Title: {slide.shapes.title.text.strip()}")
                
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
                if shape.has_table:
                    table_text = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        table_text.append(" | ".join(row_cells))
                    slide_text.append("[Table]:\n" + "\n".join(table_text))
            content.append("\n".join(slide_text))
        
        with open(out_txt_path, "w", encoding="utf-8") as f:
            f.write("\n\n".join(content))
        print(f"Successfully extracted PPTX text to {out_txt_path}")
    except Exception as e:
        print(f"Error extracting PPTX: {e}")

def extract_excel(excel_path, out_json_path, out_data_path):
    print(f"Extracting Excel: {excel_path}...")
    try:
        xl = pd.ExcelFile(excel_path)
        sheets_summary = {}
        sheets_data = {}
        
        for sheet_name in xl.sheet_names:
            df = xl.parse(sheet_name)
            
            # Clean column names
            df.columns = [str(c).strip() for c in df.columns]
            
            # Format date columns to string for JSON serialization
            for col in df.columns:
                if pd.api.types.is_datetime64_any_dtype(df[col]):
                    df[col] = df[col].dt.strftime('%Y-%m-%d')
            
            # Replace NaN with None for JSON
            df_clean = df.where(pd.notnull(df), None)
            
            cols = list(df.columns)
            num_rows = len(df)
            
            sample = df_clean.head(10).to_dict(orient="records")
            dtypes = {col: str(dtype) for col, dtype in df.dtypes.items()}
            
            stats = {}
            for col in df.columns:
                stats[col] = {
                    "unique_count": int(df[col].nunique()),
                    "null_count": int(df[col].isnull().sum()),
                    "sample_values": list(df[col].dropna().unique()[:10])
                }
                stats[col]["sample_values"] = [
                    str(x) if isinstance(x, (pd.Timestamp, str)) else (int(x) if isinstance(x, (int, pd.Int64Dtype)) else float(x) if isinstance(x, float) else x)
                    for x in stats[col]["sample_values"]
                ]
                
            sheets_summary[sheet_name] = {
                "columns": cols,
                "row_count": num_rows,
                "dtypes": dtypes,
                "stats": stats,
                "sample": sample
            }
            
            # Save full sheet records
            sheets_data[sheet_name] = df_clean.to_dict(orient="records")
            
        with open(out_json_path, "w", encoding="utf-8") as f:
            json.dump(sheets_summary, f, indent=4, default=str)
        print(f"Successfully wrote Excel summary to {out_json_path}")
        
        # Save full data as data.js for use in dashboard directly
        with open(out_data_path, "w", encoding="utf-8") as f:
            f.write("const DASHBOARD_DATA = ")
            json.dump(sheets_data, f, indent=2, default=str)
            f.write(";\n")
        print(f"Successfully wrote Excel full data to {out_data_path}")
        
    except Exception as e:
        print(f"Error extracting Excel: {e}")

if __name__ == "__main__":
    # Target files in workspace
    workspace_dir = os.path.dirname(os.path.abspath(__file__))
    
    pptx_file = os.path.join(workspace_dir, "Audience_Reach_Analysis_Presentation.pptx")
    excel_file = os.path.join(workspace_dir, "Book1.xlsx")
    
    out_txt_path = os.path.join(workspace_dir, "pptx_content.txt")
    out_json_path = os.path.join(workspace_dir, "excel_summary.json")
    out_data_path = os.path.join(workspace_dir, "data.js")
    
    if os.path.exists(pptx_file):
        extract_pptx(pptx_file, out_txt_path)
    else:
        print(f"PPTX file not found at {pptx_file}")
        
    if os.path.exists(excel_file):
        extract_excel(excel_file, out_json_path, out_data_path)
    else:
        print(f"Excel file not found at {excel_file}")
        
    print("\nExtraction complete! Files created in P5-barc dashboard -A:")
    print(" - pptx_content.txt (PowerPoint Text Content)")
    print(" - excel_summary.json (Excel Columns and Metadata)")
    print(" - data.js (JavaScript data object containing all Excel rows)")
