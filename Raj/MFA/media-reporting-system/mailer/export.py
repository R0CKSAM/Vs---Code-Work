"""
Media Reporting System - Mailer Data Export Engine
Exports mailer articles and advertisements to CSV and Excel (openpyxl) formats.
"""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


class MailerExporter:
    """Exports structured mailer intelligence into Excel and CSV files."""

    def __init__(self, export_dir: Optional[Path] = None):
        if export_dir:
            self.export_dir = Path(export_dir)
        else:
            _here = Path(__file__).resolve().parent
            _root = _here.parent
            self.export_dir = _root / "storage" / "reports"
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def export_csv(
        self,
        mailer_data: Dict[str, Any],
        filename: Optional[str] = None,
    ) -> Path:
        """Exports mailer articles and advertisements to CSV format."""
        if not filename:
            now_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            filename = f"mailer_report_{now_str}.csv"

        out_path = self.export_dir / filename
        date_str = mailer_data.get("received_datetime") or mailer_data.get("created_at") or ""
        source = mailer_data.get("source_name") or "Best Media Info"
        heading = mailer_data.get("newsletter_heading") or ""

        with open(out_path, mode="w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)

            # Section 1: Editorial Articles
            writer.writerow(["=== EDITORIAL ARTICLES ==="])
            writer.writerow([
                "Date",
                "Source",
                "Newsletter Heading",
                "Article Title",
                "Description",
                "Source URL",
                "Image URL",
            ])
            for art in mailer_data.get("articles", []):
                writer.writerow([
                    date_str,
                    art.get("source_name") or source,
                    heading,
                    art.get("title", ""),
                    art.get("description", ""),
                    art.get("source_url", ""),
                    art.get("image_url", ""),
                ])

            writer.writerow([])
            # Section 2: Advertisements
            writer.writerow(["=== ADVERTISEMENTS ==="])
            writer.writerow([
                "Date",
                "Source",
                "Advertiser",
                "Advertisement",
                "Target URL",
                "Image URL",
            ])
            for ad in mailer_data.get("advertisements", []):
                writer.writerow([
                    date_str,
                    source,
                    ad.get("advertiser", ""),
                    ad.get("title", ""),
                    ad.get("target_url", ""),
                    ad.get("image_url", ""),
                ])

        return out_path

    def export_excel(
        self,
        mailer_data: Dict[str, Any],
        filename: Optional[str] = None,
    ) -> Path:
        """Exports mailer data to a styled multi-tab Excel workbook."""
        if not filename:
            now_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            filename = f"mailer_report_{now_str}.xlsx"

        out_path = self.export_dir / filename
        wb = openpyxl.Workbook()

        date_str = mailer_data.get("received_datetime") or mailer_data.get("created_at") or ""
        source = mailer_data.get("source_name") or "Best Media Info"
        heading = mailer_data.get("newsletter_heading") or ""

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        ad_header_fill = PatternFill(start_color="7C2D12", end_color="7C2D12", fill_type="solid")

        # Tab 1: Articles
        ws_articles = wb.active
        ws_articles.title = "Articles"

        article_headers = [
            "Date", "Source", "Newsletter Heading", "Article Title",
            "Description", "Source URL", "Image URL"
        ]
        ws_articles.append(article_headers)
        for col_idx in range(1, len(article_headers) + 1):
            cell = ws_articles.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for art in mailer_data.get("articles", []):
            ws_articles.append([
                date_str,
                art.get("source_name") or source,
                heading,
                art.get("title", ""),
                art.get("description", ""),
                art.get("source_url", ""),
                art.get("image_url", ""),
            ])

        # Adjust column widths
        for col in ws_articles.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws_articles.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

        # Tab 2: Advertisements
        ws_ads = wb.create_sheet(title="Advertisements")
        ad_headers = [
            "Date", "Source", "Advertiser", "Advertisement",
            "Target URL", "Image URL"
        ]
        ws_ads.append(ad_headers)
        for col_idx in range(1, len(ad_headers) + 1):
            cell = ws_ads.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = ad_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for ad in mailer_data.get("advertisements", []):
            ws_ads.append([
                date_str,
                source,
                ad.get("advertiser", ""),
                ad.get("title", ""),
                ad.get("target_url", ""),
                ad.get("image_url", ""),
            ])

        for col in ws_ads.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws_ads.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

        wb.save(out_path)
        return out_path


mailer_exporter = MailerExporter()
