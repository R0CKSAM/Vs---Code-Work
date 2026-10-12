from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from market_mapping import normalize_city_name, normalize_headend_name, normalize_market_name


CLOSED_REMARKS = ("CLOSE", "CLOSURE", "CLOSED", "DELETE", "DELETION", "REMOVE", "REMOVAL", "OFF", "DISCONTINU")
AVAILABLE_REMARKS = ("ADD", "ADDITION", "OPEN", "ACTIVE", "AVAILABLE", "REOPEN")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", _text(value).upper())


def _key(value: Any) -> str:
    return _text(value).upper()


def _week_label(week: int, year: int) -> str:
    return f"Wk-{week:02d}'{str(year)[-2:]}"


def parse_status_week(path: Path, worksheet: Any | None = None) -> str | None:
    match = re.search(r"WK[-\s_]*(\d{1,2})\D+(20\d{2})", path.stem, re.IGNORECASE)
    if match:
        return _week_label(int(match.group(1)), int(match.group(2)))

    if worksheet is not None:
        for row in worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True):
            for value in row:
                match = re.search(r"WK[-\s_]*(\d{1,2})\D+(20\d{2})", _text(value), re.IGNORECASE)
                if match:
                    return _week_label(int(match.group(1)), int(match.group(2)))
    return None


def classify_remark(value: Any) -> str:
    remark = _key(value)
    if any(token in remark for token in CLOSED_REMARKS):
        return "closed"
    if any(token in remark for token in AVAILABLE_REMARKS):
        return "available"
    return "unknown"


@dataclass(frozen=True)
class HeadendEvent:
    week: str
    transmission: str
    market: str
    city: str
    network: str
    crn_no: str
    remark: str
    status: str
    source_file: str


@dataclass
class HeadendStatusData:
    available_weeks: set[str]
    events: list[HeadendEvent]

    def __post_init__(self) -> None:
        self._by_crn: dict[tuple[str, str], HeadendEvent] = {}
        self._by_location: dict[tuple[str, str, str, str, str], HeadendEvent] = {}
        for event in self.events:
            if event.crn_no:
                self._by_crn[(event.week, _key(event.crn_no))] = event
            location_key = (
                event.week,
                _key(event.transmission),
                _key(event.market),
                _key(event.city),
                _key(event.network),
            )
            self._by_location[location_key] = event

    def match(self, record: dict[str, Any], week: str) -> HeadendEvent | None:
        crn_no = _key(record.get("crn_no"))
        if crn_no and (week, crn_no) in self._by_crn:
            return self._by_crn[(week, crn_no)]
        location_key = (
            week,
            _key(record.get("transmission")),
            _key(normalize_market_name(_text(record.get("market")))),
            _key(normalize_city_name(_text(record.get("city")))),
            _key(normalize_headend_name(_text(record.get("head_end")))),
        )
        return self._by_location.get(location_key)


def read_status_workbook(path: Path) -> tuple[str, list[HeadendEvent]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Networks Details"] if "Networks Details" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        week = parse_status_week(path, worksheet)
        if not week:
            raise ValueError(f"Could not determine week from {path.name}")

        header_row = None
        header_map: dict[str, int] = {}
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(30, worksheet.max_row), values_only=True),
            start=1,
        ):
            keys = [_header(value) for value in row]
            if "NETWORK" in keys and "REMARKS" in keys:
                header_row = row_number
                header_map = {key: index for index, key in enumerate(keys) if key}
                break
        if header_row is None:
            raise ValueError(f"Could not find NETWORK/REMARKS headers in {path.name}")

        events: list[HeadendEvent] = []
        for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            def cell(*names: str) -> Any:
                index = next((header_map[name] for name in names if name in header_map), None)
                return values[index] if index is not None and index < len(values) else None

            network = normalize_headend_name(_text(cell("NETWORK", "HEADEND", "HEADENDNAME")))
            crn_no = _text(cell("CRNNO", "CRN"))
            remark = _text(cell("REMARKS", "REMARK", "STATUS"))
            if not network and not crn_no:
                continue
            events.append(
                HeadendEvent(
                    week=week,
                    transmission=_text(cell("TRANSMISSION")),
                    market=normalize_market_name(_text(cell("MARKET"))),
                    city=normalize_city_name(_text(cell("CITY"))),
                    network=network,
                    crn_no=crn_no,
                    remark=remark,
                    status=classify_remark(remark),
                    source_file=path.name,
                )
            )
        return week, events
    finally:
        workbook.close()


def load_headend_status_data(folder: Path) -> HeadendStatusData:
    available_weeks: set[str] = set()
    events: list[HeadendEvent] = []
    if not folder.exists():
        return HeadendStatusData(available_weeks, events)
    for path in sorted(folder.glob("*.xlsx"), key=lambda item: item.name.lower()):
        if path.name.startswith("~$"):
            continue
        week, workbook_events = read_status_workbook(path)
        available_weeks.add(week)
        events.extend(workbook_events)
    return HeadendStatusData(available_weeks, events)

